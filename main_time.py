# call all over functions, make everything look nice, ...
from zero_crossings import zero_crossings
from total_variation import total_variation
from error_pdf import error_pdf
from RMS_error import RMS_error
from RMS_u import RMS_input
from derivative_e_RMS import RMS_DERe
from derivative_u_RMS import RMS_DERu
from contributions import contributions
import pandas as pd
from openpyxl.styles import PatternFill
from box_plot import generate_plots

pd.set_option('display.max_columns', None)
pd.set_option('display.width', None)

'''Output an Excel file'''
highlight = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

with pd.ExcelWriter("results.xlsx", engine="openpyxl") as writer:
    # === Zero-crossings of e ===
    data, stat_res, array = zero_crossings()

    data.to_excel(writer, sheet_name="ZeroCrossings_data")
    stat_res.to_excel(writer, sheet_name="ZeroCrossings_stats")

    sheet = writer.sheets["ZeroCrossings_stats"]

    # Highlight p-values (first column)
    for row in sheet.iter_rows(min_row=2):
        cell = row[1]
        if float(cell.value) < 0.05:
            cell.fill = highlight

    ar = [[array[:, 0], array[:, 3]],[array[:, 1], array[:, 4]],[array[:, 2], array[:, 5]]]
    generate_plots(ar, stat_res["p_val"].values, stat_res["effect_size"].values, ["-", "-", "-"], ["Position ZC", "Velocity ZC", "Acceleration ZC"])

    # === Total Variation of u ===
    data2, stat_res2, array2 = total_variation()

    data2.to_excel(writer, sheet_name="TotalVariation_data")
    stat_res2.to_excel(writer, sheet_name="TotalVariation_stats")

    sheet = writer.sheets["TotalVariation_stats"]

    for row in sheet.iter_rows(min_row=2):
        cell = row[1]
        if float(cell.value) < 0.05:
            cell.fill = highlight

    ar2 = [[array2[:, 0], array2[:, 3]], [array2[:, 1], array2[:, 4]], [array2[:, 2], array2[:, 5]]]
    generate_plots(ar2, stat_res2["p_val"].values, stat_res2["effect_size"].values, ["-", "-", "-"],
                   ["Position TV", "Velocity TV", "Acceleration TV"])

    # === Mean and 1-sigma Interval of e ===
    data3, stat_res3, data4, stat_res4, array3, array4 = error_pdf()

    data4.to_excel(writer, sheet_name="ErrorPDF_data")
    stat_res4.to_excel(writer, sheet_name="ErrorPDF_stats")

    sheet = writer.sheets["ErrorPDF_stats"]

    for row in sheet.iter_rows(min_row=2):
        cell = row[1]
        if float(cell.value) < 0.05:
            cell.fill = highlight

    ar4 = [[array4[:, 0], array4[:, 3]], [array4[:, 1], array4[:, 4]], [array4[:, 2], array4[:, 5]]]
    generate_plots(ar4, stat_res4["p_val"].values, stat_res4["effect_size"].values, ["deg", "deg", "deg"],
                   ["Position Interval", "Velocity Interval", "Acceleration Interval"])

    # === RMS of e ===
    data5, stat_res5, array5 = RMS_error()

    data5.to_excel(writer, sheet_name="RMSe_data")
    stat_res5.to_excel(writer, sheet_name="RMSe_stats")

    sheet = writer.sheets["RMSe_stats"]

    for row in sheet.iter_rows(min_row=2):
        cell = row[1]
        if float(cell.value) < 0.05:
            cell.fill = highlight

    ar5 = [[array5[:, 0], array5[:, 3]], [array5[:, 1], array5[:, 4]], [array5[:, 2], array5[:, 5]]]
    generate_plots(ar5, stat_res5["p_val"].values, stat_res5["effect_size"].values, ["deg", "deg", "deg"],
                   ["Position e", "Velocity e", "Acceleration e"])

    # === RMS of u ===
    data6, stat_res6, array6 = RMS_input()

    data6.to_excel(writer, sheet_name="RMSu_data")
    stat_res6.to_excel(writer, sheet_name="RMSu_stats")

    sheet = writer.sheets["RMSu_stats"]

    for row in sheet.iter_rows(min_row=2):
        cell = row[1]
        if float(cell.value) < 0.05:
            cell.fill = highlight

    ar6 = [[array6[:, 0], array6[:, 3]], [array6[:, 1], array6[:, 4]], [array6[:, 2], array6[:, 5]]]
    generate_plots(ar6, stat_res6["p_val"].values, stat_res6["effect_size"].values, ["-", "-", "-"],
                   ["Position u", "Velocity u", "Acceleration u"])

    # === RMS of DERe ===
    data7, stat_res7, array7 = RMS_DERe()

    data7.to_excel(writer, sheet_name="RMSDERe_data")
    stat_res7.to_excel(writer, sheet_name="RMSDERe_stats")

    sheet = writer.sheets["RMSDERe_stats"]

    for row in sheet.iter_rows(min_row=2):
        cell = row[1]
        if float(cell.value) < 0.05:
            cell.fill = highlight

    ar7 = [[array7[:, 0], array7[:, 3]], [array7[:, 1], array7[:, 4]], [array7[:, 2], array7[:, 5]]]
    generate_plots(ar7, stat_res7["p_val"].values, stat_res7["effect_size"].values, ["deg/s", "deg/s", "deg/s"],
                   ["Position e'", "Velocity e'", "Acceleration e'"])

    # === RMS of DERu ===
    data8, stat_res8, array8 = RMS_DERu()

    data8.to_excel(writer, sheet_name="RMSDERu_data")
    stat_res8.to_excel(writer, sheet_name="RMSDERu_stats")

    sheet = writer.sheets["RMSDERu_stats"]

    for row in sheet.iter_rows(min_row=2):
        cell = row[1]
        if float(cell.value) < 0.05:
            cell.fill = highlight

    ar8 = [[array8[:, 0], array8[:, 3]], [array8[:, 1], array8[:, 4]], [array8[:, 2], array8[:, 5]]]
    generate_plots(ar8, stat_res8["p_val"].values, stat_res8["effect_size"].values, ["/s", "/s", "/s"],
                   ["Position u'", "Velocity u'", "Acceleration u'"])

    # === Contributions to u ===
    data9, stat_res9, data10, stat_res10, data11, stat_res11, array9, array10, array11 = contributions()
    ar9 = [[array9[:, 0], array9[:, 3]], [array9[:, 1], array9[:, 4]], [array9[:, 2], array9[:, 5]]]
    generate_plots(ar9, stat_res9["p_val"].values, stat_res9["effect_size"].values, ["pct", "pct", "pct"],
                   ["Position cont d", "Velocity cont d", "Acceleration cont d"])
    ar10 = [[array10[:, 0], array10[:, 3]], [array10[:, 1], array10[:, 4]], [array10[:, 2], array10[:, 5]]]
    generate_plots(ar10, stat_res10["p_val"].values, stat_res10["effect_size"].values, ["pct", "pct", "pct"],
                   ["Position cont t", "Velocity cont t", "Acceleration cont t"])
    ar11 = [[array11[:, 0], array11[:, 3]], [array11[:, 1], array11[:, 4]], [array11[:, 2], array11[:, 5]]]
    generate_plots(ar11, stat_res11["p_val"].values, stat_res11["effect_size"].values, ["pct", "pct", "pct"],
                   ["Position cont n", "Velocity cont n", "Acceleration cont n"])

    # --- WRITE DATA SHEET ---
    start_row = 0

    # f_d contribution
    data9.to_excel(writer, sheet_name="Contrib_data", startrow=start_row)
    start_row += len(data9) + 3

    # f_t contribution
    data10.to_excel(writer, sheet_name="Contrib_data", startrow=start_row)
    start_row += len(data10) + 3

    # noise contribution
    data11.to_excel(writer, sheet_name="Contrib_data", startrow=start_row)

    # --- WRITE STATS SHEET ---
    start_row = 0

    # f_d stats
    stat_res9.to_excel(writer, sheet_name="Contrib_stats", startrow=start_row)
    start_row += len(stat_res9) + 3

    # f_t stats
    stat_res10.to_excel(writer, sheet_name="Contrib_stats", startrow=start_row)
    start_row += len(stat_res10) + 3

    # noise stats
    stat_res11.to_excel(writer, sheet_name="Contrib_stats", startrow=start_row)

    # --- APPLY HIGHLIGHTING ---
    sheet = writer.sheets["Contrib_stats"]

    current_row = 0
    for stat_res in [stat_res9, stat_res10, stat_res11]:

        for row in sheet.iter_rows(min_row=current_row + 2,
                                   max_row=current_row + 1 + len(stat_res)):

            cell = row[1]  # p-value column
            if float(cell.value) < 0.05:
                cell.fill = highlight

        current_row += len(stat_res) + 3

'''If a .txt file with the outputs are more useful, use this
with open("results.txt", "w") as f:
    # Zero-crossings of e
    f.write("=== Zero-crossings of e ===\n")
    data, stat_res = zero_crossings()

    f.write("\nMean Zero-crossings:\n")
    f.write(data.to_string())
    f.write("\n\nStatistical Results:\n")
    f.write(stat_res.to_string())
    f.write("\n\n")

    # Total Variation of u
    f.write("=== Total Variation of u ===\n")
    data2, stat_res2 = total_variation()

    f.write("\nMean Total Variation:\n")
    f.write(data2.to_string())
    f.write("\n\nStatistical Results:\n")
    f.write(stat_res2.to_string())
    f.write("\n\n")

    # Mean and 1-sigma Interval of e
    f.write("=== Mean and 1-sigma Interval of e ===\n")
    data3, stat_res3, data4, stat_res4 = error_pdf()

    f.write("\nMean 1-sigma Interval Width:\n")
    f.write(data4.to_string())
    f.write("\n\nStatistical Results:\n")
    f.write(stat_res4.to_string())
    f.write("\n")
'''